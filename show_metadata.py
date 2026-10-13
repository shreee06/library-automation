"""
show_metadata.py

Prints every row of output/email_metadata.xlsx as a readable
"column: value" block, one block per email. Read-only.

Usage:
    python show_metadata.py
    python show_metadata.py path/to/other.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_PATH = BASE_DIR / "output" / "email_metadata.xlsx"
ATTACHMENT_DIR = BASE_DIR / "email_input"


def show(path: Path) -> None:
    ws = load_workbook(path, read_only=True)["Emails"]
    rows = list(ws.iter_rows(values_only=True))
    headers, data = rows[0], rows[1:]
    width = max(len(h) for h in headers)

    print(f"{path}  --  {len(data)} row(s)\n")
    for n, row in enumerate(data, start=1):
        print(f"{'=' * 20} Row {n} of {len(data)} {'=' * 20}")
        for header, value in zip(headers, row):
            text = "(blank)" if value in (None, "") else str(value)
            lines = text.splitlines() or [text]
            # Collapse runs of blank lines (HTML-derived bodies have many).
            lines = [ln.rstrip() for ln in lines]
            compact = [ln for i, ln in enumerate(lines) if ln or (i and lines[i - 1])]
            print(f"  {header + ':':<{width + 1}} {compact[0] if compact else ''}")
            for extra in compact[1:]:
                print(f"  {'':<{width + 1}} {extra}")

        print_attachment_paths(dict(zip(headers, row)).get("Attachments"), width)
        print()


def print_attachment_paths(attachments: str | None, width: int) -> None:
    """Show the resolved on-disk path of each saved attachment."""
    label = "Attachment Paths:"
    names = [n.strip() for n in (attachments or "").split(";") if n.strip()]
    if not names:
        print(f"  {label:<{width + 1}} (none)")
        return
    for i, name in enumerate(names):
        path = (ATTACHMENT_DIR / name).resolve()
        note = "" if path.exists() else "  [MISSING ON DISK]"
        print(f"  {label if i == 0 else '':<{width + 1}} {path}{note}")


if __name__ == "__main__":
    show(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_PATH)
