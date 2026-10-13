"""
excel_logger.py

Appends one row per checked submission to a running .xlsx master log.
Creates the file with headers if it doesn't exist yet.
"""

from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = [
    "Timestamp",
    "Student Name",
    "Registration Number",
    "Department",
    "Source File",
    "AI %",
    "Human %",
    "AI Tier",
    "Plagiarism Detected",
    "Plagiarism Match Count",
    "Status",
]

SAVE_ATTEMPTS = 3
SAVE_RETRY_SECONDS = 3


class LogWriteError(RuntimeError):
    """Raised when the log file can't be written (e.g. it's open in Excel)."""


def _ensure_workbook(log_path: Path) -> None:
    if log_path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Submissions"
    ws.append(HEADERS)
    # Light formatting: bold header row
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wb.save(log_path)


def append_row(
    log_path: str | Path,
    *,
    student_name: str,
    reg_no: str,
    department: str,
    source_file: str,
    ai_percent: float | None,
    human_percent: float | None,
    ai_tier: str | None,
    plagiarism_detected: bool | None,
    plagiarism_match_count: int | None,
    status: str,
) -> None:
    """
    Append one row to the master log. Creates the file with headers on
    first use. Retries a few times if the file is locked (open in Excel),
    then raises LogWriteError.
    """
    path = Path(log_path)
    row = [
        datetime.now().isoformat(timespec="seconds"),
        student_name,
        reg_no,
        department,
        source_file,
        ai_percent,
        human_percent,
        ai_tier,
        plagiarism_detected,
        plagiarism_match_count,
        status,
    ]

    for attempt in range(1, SAVE_ATTEMPTS + 1):
        try:
            _ensure_workbook(path)
            wb = load_workbook(path)
            ws = wb["Submissions"]
            ws.append(row)
            # openpyxl turns any string starting with "=" into a formula.
            # Filenames are untrusted input, so force text cells.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
            wb.save(path)
            return
        except PermissionError:
            if attempt < SAVE_ATTEMPTS:
                print(
                    f"  [!] '{path.name}' is locked (open in Excel?). "
                    f"Retrying in {SAVE_RETRY_SECONDS}s "
                    f"({attempt}/{SAVE_ATTEMPTS})..."
                )
                time.sleep(SAVE_RETRY_SECONDS)

    raise LogWriteError(
        f"Could not write to '{path}' after {SAVE_ATTEMPTS} attempts. "
        f"Close it in Excel and run again; unprocessed PDFs are still in input/."
    )


def print_log_table(log_path: str | Path) -> None:
    """Print the master log to the terminal as a database-style table."""
    path = Path(log_path)
    if not path.exists():
        print("(no log yet)")
        return

    wb = load_workbook(path, read_only=True)
    rows = [
        ["" if v is None else str(v) for v in row]
        for row in wb["Submissions"].iter_rows(min_row=2, values_only=True)
    ]
    wb.close()

    widths = [
        max(len(h), *(len(r[i]) for r in rows)) if rows else len(h)
        for i, h in enumerate(HEADERS)
    ]
    border = "+" + "+".join("-" * (w + 2) for w in widths) + "+"

    def fmt(cells: list[str]) -> str:
        return "| " + " | ".join(c.ljust(w) for c, w in zip(cells, widths)) + " |"

    print(border)
    print(fmt(HEADERS))
    print(border)
    for r in rows:
        print(fmt(r))
    print(border)
    print(f"{len(rows)} row(s) in set")


if __name__ == "__main__":
    # Quick manual test: python excel_logger.py
    append_row(
        "test_log.xlsx",
        student_name="Test Student",
        reg_no="99TEST01",
        department="TEST",
        source_file="test.pdf",
        ai_percent=12.5,
        human_percent=87.5,
        ai_tier="Fully Human",
        plagiarism_detected=False,
        plagiarism_match_count=0,
        status="Processed",
    )
    print("Wrote a test row to test_log.xlsx")
